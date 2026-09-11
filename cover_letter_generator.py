import os
import json
import re
import requests
from docx import Document
from docx.shared import Inches, Pt


# ============================================================
# CONFIGURATION
# ============================================================

BASE_DIR = os.path.dirname(os.path.abspath(__file__))

BASE_RESUME = os.path.join(
    BASE_DIR,
    "base_java_resume.docx"
)

OUTPUT_DIR = os.path.join(
    BASE_DIR,
    "generated_cover_letters"
)

LM_STUDIO_URL = "http://localhost:1234/api/v1/chat"

LM_STUDIO_MODEL = "google/gemma-4-e4b"

os.makedirs(OUTPUT_DIR, exist_ok=True)


# ============================================================
# READ BASE RESUME
# ============================================================

def read_resume(file_path):

    if not os.path.exists(file_path):
        raise FileNotFoundError(
            f"Resume not found:\n{file_path}"
        )

    document = Document(file_path)

    lines = []

    for paragraph in document.paragraphs:

        text = paragraph.text.strip()

        if text:
            lines.append(text)

    return "\n".join(lines)


# ============================================================
# GET CANDIDATE DETAILS
# ============================================================

def get_candidate_details(file_path):

    document = Document(file_path)

    lines = []

    for paragraph in document.paragraphs:

        text = paragraph.text.strip()

        if text:
            lines.append(text)

        if len(lines) >= 5:
            break

    return lines


# ============================================================
# EXTRACT AI RESPONSE
# ============================================================

def extract_ai_text(data):

    output = data.get("output")

    if isinstance(output, str):
        return output.strip()

    texts = []

    if isinstance(output, list):

        for item in output:

            if not isinstance(item, dict):
                continue

            content = item.get("content")

            if isinstance(content, str):

                texts.append(content)

            elif isinstance(content, list):

                for part in content:

                    if not isinstance(part, dict):
                        continue

                    if part.get("type") in (
                        "text",
                        "output_text"
                    ):

                        text = part.get("text", "")

                        if text:
                            texts.append(text)

    return "\n".join(texts).strip()


# ============================================================
# EXTRACT JSON
# ============================================================

def extract_json(text):

    text = text.strip()

    text = re.sub(
        r"```json",
        "",
        text,
        flags=re.IGNORECASE
    )

    text = text.replace(
        "```",
        ""
    ).strip()

    start = text.find("{")

    if start == -1:
        raise ValueError(
            "AI response does not contain JSON."
        )

    decoder = json.JSONDecoder()

    result, _ = decoder.raw_decode(
        text[start:]
    )

    return result


# ============================================================
# GENERATE COVER LETTER
# ============================================================

def generate_cover_letter(
    job_title,
    company,
    job_description,
    resume_text
):

    # Keep request reasonably small
    resume_text = resume_text[:16000]

    job_description = job_description[:7000]

    prompt = f"""
Write a professional cover letter for the candidate below.

TARGET JOB:
Job Title: {job_title}
Company: {company}

JOB DESCRIPTION:
{job_description}

CANDIDATE RESUME:
{resume_text}

STRICT RULES:

1. Use ONLY facts from the candidate resume.
2. Never invent skills, technologies, experience,
   achievements, employers, dates, education,
   certifications or projects.
3. Never exaggerate the candidate's experience.
4. Keep the candidate's fresher/entry-level status honest.
5. Mention only skills that are actually present in
   the candidate resume.
6. Do not claim the candidate has professional experience
   with a technology just because it appears in the job
   description.
7. Do not mention information that is not supported by
   the resume.
8. Make the letter specific to the target job.
9. Keep it concise and professional.
10. Do not use placeholders such as [Name] or [Company].
11. Do not include a subject line.
12. Return ONLY JSON.

Use exactly this JSON format:

{{
  "greeting": "Dear Hiring Manager,",
  "opening": "Short opening paragraph showing interest in the role.",
  "body": "One or two paragraphs connecting the candidate's relevant skills, training, projects and experience to the job.",
  "closing": "Short professional closing paragraph.",
  "sign_off": "Sincerely"
}}
"""

    payload = {
        "model": LM_STUDIO_MODEL,
        "input": prompt,
        "temperature": 0,
        "max_output_tokens": 800,
        "reasoning": "off",
        "store": False
    }

    print("\nSending cover letter request to Gemma...")

    response = requests.post(
        LM_STUDIO_URL,
        json=payload,
        timeout=300
    )

    response.raise_for_status()

    data = response.json()

    ai_text = extract_ai_text(data)

    if not ai_text:
        raise ValueError(
            "Gemma returned an empty response."
        )

    print("AI response received.")

    result = extract_json(ai_text)

    required_fields = [
        "greeting",
        "opening",
        "body",
        "closing",
        "sign_off"
    ]

    for field in required_fields:

        if field not in result:
            raise ValueError(
                f"Missing AI field: {field}"
            )

    return result


# ============================================================
# SAFE FILE NAME
# ============================================================

def safe_filename(text):

    text = str(text)

    text = re.sub(
        r'[<>:"/\\|?*]',
        "_",
        text
    )

    text = re.sub(
        r"\s+",
        "_",
        text
    )

    return text[:60]


# ============================================================
# CREATE COVER LETTER DOCX
# ============================================================

def create_cover_letter(
    data,
    candidate_details,
    job_title,
    company
):

    document = Document()

    # --------------------------------------------------------
    # PAGE SETTINGS
    # --------------------------------------------------------

    section = document.sections[0]

    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.85)

    # --------------------------------------------------------
    # DEFAULT FONT
    # --------------------------------------------------------

    document.styles["Normal"].font.name = "Calibri"
    document.styles["Normal"].font.size = Pt(11)

    # --------------------------------------------------------
    # CANDIDATE DETAILS
    # --------------------------------------------------------

    if candidate_details:

        name = document.add_paragraph()

        run = name.add_run(
            candidate_details[0]
        )

        run.bold = True
        run.font.size = Pt(15)

        for line in candidate_details[1:]:

            paragraph = document.add_paragraph()

            paragraph.paragraph_format.space_after = Pt(1)

            run = paragraph.add_run(line)

            run.font.size = Pt(9)

    # --------------------------------------------------------
    # DATE
    # --------------------------------------------------------

    document.add_paragraph("")

    date_paragraph = document.add_paragraph()

    date_paragraph.add_run(
        "Date: "
    ).bold = True

    from datetime import datetime

    date_paragraph.add_run(
        datetime.now().strftime("%d %B %Y")
    )

    # --------------------------------------------------------
    # RECIPIENT
    # --------------------------------------------------------

    document.add_paragraph("")

    recipient = document.add_paragraph()

    recipient.add_run(
        "Hiring Manager"
    ).bold = True

    document.add_paragraph(
        company
    )

    # --------------------------------------------------------
    # SUBJECT
    # --------------------------------------------------------

    subject = document.add_paragraph()

    subject.paragraph_format.space_before = Pt(8)

    subject_run = subject.add_run(
        f"Application for {job_title}"
    )

    subject_run.bold = True

    # --------------------------------------------------------
    # GREETING
    # --------------------------------------------------------

    document.add_paragraph("")

    greeting = document.add_paragraph()

    greeting.add_run(
        data["greeting"]
    )

    # --------------------------------------------------------
    # OPENING
    # --------------------------------------------------------

    document.add_paragraph("")

    opening = document.add_paragraph(
        data["opening"]
    )

    opening.paragraph_format.space_after = Pt(8)

    # --------------------------------------------------------
    # BODY
    # --------------------------------------------------------

    body = data["body"]

    # If AI returns multiple paragraphs,
    # preserve them.

    body_paragraphs = body.split("\n")

    for paragraph_text in body_paragraphs:

        paragraph_text = paragraph_text.strip()

        if not paragraph_text:
            continue

        paragraph = document.add_paragraph(
            paragraph_text
        )

        paragraph.paragraph_format.space_after = Pt(8)

    # --------------------------------------------------------
    # CLOSING
    # --------------------------------------------------------

    closing = document.add_paragraph(
        data["closing"]
    )

    closing.paragraph_format.space_after = Pt(10)

    # --------------------------------------------------------
    # SIGN OFF
    # --------------------------------------------------------

    signoff = document.add_paragraph()

    signoff.add_run(
        data["sign_off"]
    )

    document.add_paragraph("")

    # Candidate name
    if candidate_details:

        name = document.add_paragraph()

        run = name.add_run(
            candidate_details[0]
        )

        run.bold = True

    # --------------------------------------------------------
    # SAVE FILE
    # --------------------------------------------------------

    filename = (
        "Cover_Letter_"
        + safe_filename(company)
        + "_"
        + safe_filename(job_title)
        + ".docx"
    )

    output_path = os.path.join(
        OUTPUT_DIR,
        filename
    )

    document.save(output_path)

    return output_path


# ============================================================
# MAIN FUNCTION
# ============================================================

def customize_cover_letter(
    job_title,
    company,
    job_description
):

    print("\n==========================================")
    print("       AI COVER LETTER GENERATION")
    print("==========================================")

    print("\nReading base resume...")

    resume_text = read_resume(
        BASE_RESUME
    )

    candidate_details = get_candidate_details(
        BASE_RESUME
    )

    print("Base resume loaded successfully.")

    result = generate_cover_letter(
        job_title,
        company,
        job_description,
        resume_text
    )

    print("Creating cover letter...")

    output_file = create_cover_letter(
        result,
        candidate_details,
        job_title,
        company
    )

    print("\n==========================================")
    print("COVER LETTER GENERATED SUCCESSFULLY")
    print("==========================================")

    print("\nFile:")
    print(output_file)

    return output_file


# ============================================================
# TEST
# ============================================================

if __name__ == "__main__":

    test_job_title = "Java Developer"

    test_company = "Demo Company"

    test_job_description = """
    We are looking for a Java Developer Fresher.

    Requirements:
    Core Java, OOP, SQL, MySQL, exception handling,
    collections, debugging and problem solving.

    Knowledge of backend development is preferred.
    """

    try:

        customize_cover_letter(
            test_job_title,
            test_company,
            test_job_description
        )

        print("\nTEST COMPLETED SUCCESSFULLY.")

    except Exception as error:

        print("\n==========================================")
        print("ERROR")
        print("==========================================")

        print(error)