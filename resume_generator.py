import os
import json
import re
import requests
from docx import Document
from docx.shared import Inches, Pt


# ============================================================
# CONFIGURATION
# ============================================================

BASE_DIR = os.path.dirname(os.path.abspath(__file__))

BASE_RESUME = os.path.join(
    BASE_DIR,
    "base_java_resume.docx"
)

OUTPUT_DIR = os.path.join(
    BASE_DIR,
    "generated_resumes"
)

LM_STUDIO_URL = "http://localhost:1234/api/v1/chat"

LM_STUDIO_MODEL = "google/gemma-4-e4b"

os.makedirs(OUTPUT_DIR, exist_ok=True)


# ============================================================
# READ RESUME
# ============================================================

def read_resume(file_path):

    if not os.path.exists(file_path):
        raise FileNotFoundError(
            f"Resume not found:\n{file_path}"
        )

    document = Document(file_path)

    lines = []

    for paragraph in document.paragraphs:

        text = paragraph.text.strip()

        if text:
            lines.append(text)

    return "\n".join(lines)


# ============================================================
# GET HEADER
# ============================================================

def get_header(file_path):

    document = Document(file_path)

    lines = []

    for paragraph in document.paragraphs:

        text = paragraph.text.strip()

        if text:
            lines.append(text)

        if len(lines) >= 5:
            break

    return lines


# ============================================================
# EXTRACT AI TEXT
# ============================================================

def extract_ai_text(data):

    output = data.get("output")

    if isinstance(output, str):
        return output.strip()

    texts = []

    if isinstance(output, list):

        for item in output:

            if not isinstance(item, dict):
                continue

            content = item.get("content")

            if isinstance(content, str):
                texts.append(content)

            elif isinstance(content, list):

                for part in content:

                    if not isinstance(part, dict):
                        continue

                    if part.get("type") in (
                        "text",
                        "output_text"
                    ):

                        text = part.get("text", "")

                        if text:
                            texts.append(text)

    return "\n".join(texts).strip()


# ============================================================
# EXTRACT JSON
# ============================================================

def extract_json(text):

    text = text.strip()

    text = re.sub(
        r"```json",
        "",
        text,
        flags=re.IGNORECASE
    )

    text = text.replace("```", "").strip()

    start = text.find("{")

    if start == -1:
        raise ValueError(
            "AI did not return JSON."
        )

    decoder = json.JSONDecoder()

    result, _ = decoder.raw_decode(
        text[start:]
    )

    return result


# ============================================================
# CONVERT TO LIST
# ============================================================

def make_list(value):

    if value is None:
        return []

    if isinstance(value, list):

        return [
            str(x).strip()
            for x in value
            if str(x).strip()
        ]

    if isinstance(value, str):

        result = []

        for line in value.splitlines():

            line = line.strip()

            line = re.sub(
                r"^[•*\-\d.)]+\s*",
                "",
                line
            )

            if line:
                result.append(line)

        return result

    return [str(value)]


# ============================================================
# AI RESUME CUSTOMIZATION
# ============================================================

def customize_with_ai(
    job_title,
    company,
    job_description,
    resume_text
):

    # --------------------------------------------------------
    # IMPORTANT:
    # Keep the prompt small so Gemma responds faster.
    # --------------------------------------------------------

    resume_text = resume_text[:9000]

    job_description = job_description[:5000]

    prompt = f"""
Customize a fresher Java resume for this job.

STRICT RULES:
1. Use ONLY information from the BASE RESUME.
2. Never invent skills, experience, companies, dates,
   projects, certifications or achievements.
3. Do not upgrade basic knowledge to advanced expertise.
4. Keep fresher experience honest.
5. Select only information relevant to the job.
6. Do not add technologies only because they are in the job.
7. Keep every item concise and ATS-friendly.
8. Return ONLY valid JSON.
9. Do not use markdown.
10. Do not explain your answer.

JOB TITLE:
{job_title}

COMPANY:
{company}

JOB DESCRIPTION:
{job_description}

BASE RESUME:
{resume_text}

Return exactly this JSON structure:

{{
"summary":"2-3 sentence professional summary",
"skills":["relevant skill 1","relevant skill 2"],
"training":["relevant training"],
"experience":["relevant experience"],
"projects":["relevant project"],
"certifications":["relevant certification"],
"education":["education"],
"objective":"short career objective"
}}
"""

    payload = {
        "model": LM_STUDIO_MODEL,
        "input": prompt,
        "temperature": 0,
        "max_output_tokens": 600,
        "reasoning": "off",
        "store": False
    }

    print("\nSending request to Gemma...")

    try:

        response = requests.post(
            LM_STUDIO_URL,
            json=payload,
            timeout=300
        )

        response.raise_for_status()

    except requests.exceptions.Timeout:

        raise RuntimeError(
            "Gemma timed out while generating the resume. "
            "The AI request took longer than 300 seconds."
        )

    except requests.exceptions.RequestException as error:

        raise RuntimeError(
            f"Gemma request failed:\n{error}"
        )

    data = response.json()

    ai_text = extract_ai_text(data)

    if not ai_text:

        raise ValueError(
            "Gemma returned an empty response."
        )

    print("AI response received.")

    result = extract_json(ai_text)

    required = [
        "summary",
        "skills",
        "training",
        "experience",
        "projects",
        "certifications",
        "education",
        "objective"
    ]

    for key in required:

        if key not in result:

            raise ValueError(
                f"Missing AI field: {key}"
            )

    return result


# ============================================================
# ADD HEADING
# ============================================================

def add_heading(document, text):

    paragraph = document.add_paragraph()

    paragraph.paragraph_format.space_before = Pt(7)
    paragraph.paragraph_format.space_after = Pt(3)

    run = paragraph.add_run(text)

    run.bold = True
    run.font.size = Pt(11)


# ============================================================
# ADD TEXT
# ============================================================

def add_text(document, text):

    paragraph = document.add_paragraph()

    paragraph.paragraph_format.space_after = Pt(3)

    run = paragraph.add_run(str(text))

    run.font.size = Pt(10)


# ============================================================
# ADD BULLET
# ============================================================

def add_bullet(document, text):

    paragraph = document.add_paragraph(
        style="List Bullet"
    )

    paragraph.paragraph_format.space_after = Pt(1)

    run = paragraph.add_run(str(text))

    run.font.size = Pt(10)


# ============================================================
# SAFE FILE NAME
# ============================================================

def safe_filename(text):

    text = str(text)

    text = re.sub(
        r'[<>:"/\\|?*]',
        "_",
        text
    )

    text = re.sub(
        r"\s+",
        "_",
        text
    )

    return text[:60]


# ============================================================
# CREATE DOCX
# ============================================================

def create_resume(
    data,
    header,
    job_title,
    company
):

    document = Document()

    section = document.sections[0]

    section.top_margin = Inches(0.55)
    section.bottom_margin = Inches(0.55)
    section.left_margin = Inches(0.65)
    section.right_margin = Inches(0.65)

    document.styles["Normal"].font.name = "Calibri"
    document.styles["Normal"].font.size = Pt(10)

    # --------------------------------------------------------
    # HEADER
    # --------------------------------------------------------

    if header:

        name = document.add_paragraph()

        name.alignment = 1

        run = name.add_run(header[0])

        run.bold = True
        run.font.size = Pt(16)

        for line in header[1:]:

            paragraph = document.add_paragraph()

            paragraph.alignment = 1
            paragraph.paragraph_format.space_after = Pt(1)

            run = paragraph.add_run(line)

            run.font.size = Pt(9)

    # --------------------------------------------------------
    # SUMMARY
    # --------------------------------------------------------

    add_heading(
        document,
        "PROFESSIONAL SUMMARY"
    )

    add_text(
        document,
        data["summary"]
    )

    # --------------------------------------------------------
    # SKILLS
    # --------------------------------------------------------

    add_heading(
        document,
        "TECHNICAL SKILLS"
    )

    for item in make_list(data["skills"]):

        add_bullet(
            document,
            item
        )

    # --------------------------------------------------------
    # JAVA TRAINING
    # --------------------------------------------------------

    training = make_list(
        data["training"]
    )

    if training:

        add_heading(
            document,
            "JAVA TRAINING"
        )

        for item in training:

            add_bullet(
                document,
                item
            )

    # --------------------------------------------------------
    # EXPERIENCE
    # --------------------------------------------------------

    experience = make_list(
        data["experience"]
    )

    if experience:

        add_heading(
            document,
            "PROFESSIONAL EXPERIENCE"
        )

        for item in experience:

            add_bullet(
                document,
                item
            )

    # --------------------------------------------------------
    # PROJECTS
    # --------------------------------------------------------

    projects = make_list(
        data["projects"]
    )

    if projects:

        add_heading(
            document,
            "RELEVANT PROJECTS"
        )

        for item in projects:

            add_bullet(
                document,
                item
            )

    # --------------------------------------------------------
    # CERTIFICATIONS
    # --------------------------------------------------------

    certifications = make_list(
        data["certifications"]
    )

    if certifications:

        add_heading(
            document,
            "TRAININGS & CERTIFICATIONS"
        )

        for item in certifications:

            add_bullet(
                document,
                item
            )

    # --------------------------------------------------------
    # EDUCATION
    # --------------------------------------------------------

    education = make_list(
        data["education"]
    )

    if education:

        add_heading(
            document,
            "EDUCATION"
        )

        for item in education:

            add_bullet(
                document,
                item
            )

    # --------------------------------------------------------
    # OBJECTIVE
    # --------------------------------------------------------

    if data.get("objective"):

        add_heading(
            document,
            "CAREER OBJECTIVE"
        )

        add_text(
            document,
            data["objective"]
        )

    # --------------------------------------------------------
    # SAVE
    # --------------------------------------------------------

    filename = (
        "Customized_Resume_"
        + safe_filename(company)
        + "_"
        + safe_filename(job_title)
        + ".docx"
    )

    output_path = os.path.join(
        OUTPUT_DIR,
        filename
    )

    document.save(output_path)

    return output_path


# ============================================================
# MAIN
# ============================================================

def customize_resume(
    job_title,
    company,
    job_description
):

    print("\n==========================================")
    print("       AI RESUME CUSTOMIZATION")
    print("==========================================")

    print("\nReading base resume...")

    resume_text = read_resume(
        BASE_RESUME
    )

    header = get_header(
        BASE_RESUME
    )

    print("Base resume loaded successfully.")

    result = customize_with_ai(
        job_title,
        company,
        job_description,
        resume_text
    )

    print("Creating customized resume...")

    output_file = create_resume(
        result,
        header,
        job_title,
        company
    )

    print("\n==========================================")
    print("RESUME GENERATED SUCCESSFULLY")
    print("==========================================")

    print("\nFile:")
    print(output_file)

    return output_file


# ============================================================
# TEST
# ============================================================

if __name__ == "__main__":

    test_job_title = "Java Developer"

    test_company = "Demo Company"

    test_job_description = """
    We are looking for a Java Developer Fresher.

    Requirements:
    Core Java, OOP, SQL, MySQL, exception handling,
    collections, debugging and problem solving.

    Knowledge of backend development is preferred.
    """

    try:

        customize_resume(
            test_job_title,
            test_company,
            test_job_description
        )

        print("\nTEST COMPLETED SUCCESSFULLY.")

    except Exception as error:

        print("\n==========================================")
        print("ERROR")
        print("==========================================")

        print(error)